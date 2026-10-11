"""Build the static site and original study downloads from data/study.json."""
import csv
import io
import json
import re
import zipfile
from collections import defaultdict
from html import escape as h
from pathlib import Path
from urllib.parse import unquote,urlparse
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
D=json.loads((ROOT/'data/study.json').read_text())
LESSONS=D['lessons']; S=D['stats']; DL=ROOT/'downloads'
(DL/'episodes').mkdir(parents=True,exist_ok=True)
TITLE='风骚律师逐集英语学习册'
SITE='https://chuanwang66.github.io/better-call-saul-english/'
REPO='https://github.com/chuanwang66/better-call-saul-english'
KINDS={'word':'关键词','phrase':'词组 / 术语'}
VERSION=f"V{D['version']}"
UPDATED=' 年 '.join(D['updated'].split('-')[:2])+' 月 '+D['updated'].split('-')[2]+' 日'
for e in LESSONS:
    assert 1<=len(e['cards'])<=15 and e['word_count']<=200
    for c in e['cards']:
        assert c['part_of_speech'] and c['definition_en']
        assert c['type']!='word' or (c['ipa'].startswith('/') and c['ipa'].endswith('/'))

def lexical_html(c):
    ipa=f'<span class="ipa" lang="en-US"><span class="meta-label">美式 IPA</span> {h(c["ipa"])}</span>' if c['ipa'] else ''
    return f'<div class="lexical-meta">{ipa}<span class="pos"><span class="meta-label">词性</span> <span lang="en">{h(c["part_of_speech"])}</span></span></div><p class="definition" lang="en"><b>English</b> {h(c["definition_en"])}</p>'

def lexical_lines(c):
    lines=[]
    if c['ipa']:lines.append('美式 IPA：'+c['ipa'])
    lines.extend(['词性：'+c['part_of_speech'],'English definition: '+c['definition_en']])
    if c['base_form']:lines.append('原形：'+c['base_form'])
    return lines

def card_html(c,n):
    search=' '.join([c['term'],c['meaning'],c['usage'],c['matched_form'],c['context'],c['definition_en'],c['part_of_speech'],c['base_form']]+[v for ex in c['examples'] for v in ex.values()]).lower()
    examples=''.join(f'<li><p class="english" lang="en">{h(ex["en"])}</p><p class="zh" lang="zh-CN">{h(ex["zh"])}</p></li>' for ex in c['examples'])
    context=f'<p class="scene zh"><b>本集用法</b>{h(c["context"])}</p>' if c['context'] else ''
    form=f'<p class="form"><b>字幕中出现</b> <span lang="en">{h(c["matched_form"])}</span></p>'
    if c['base_form']:form+=f'<p class="form"><b>原形</b> <span lang="en">{h(c["base_form"])}</span></p>'
    return f'''<article class="card" data-kind="{c['type']}" data-search="{h(search,quote=True)}">
<div class="card-top"><span class="category">{KINDS[c['type']]}</span><span class="time">回看 {c['time']}</span></div>
<h3 lang="en"><span class="number">{n:02}</span>{h(c['term'])}</h3>
{lexical_html(c)}
<p class="meaning zh">{h(c['meaning'])}</p>{form}{context}
<p class="usage zh"><b>用法</b>{h(c['usage'])}</p><ol class="examples">{examples}</ol></article>'''

def practice_html(e):
    chosen=[e['cards'][i] for i in sorted(set([0,len(e['cards'])//2,len(e['cards'])-1]))]
    assert len(chosen)==3
    items=[]
    for c in chosen:
        ex=c['examples'][1]
        items.append(f'<li><p class="zh">{h(ex["zh"])}</p><details class="answer"><summary>查看参考答案</summary><p class="english" lang="en">{h(ex["en"])}</p><p class="zh">目标词条：{h(c["term"])}</p></details></li>')
    return '<div class="practice"><h3>汉译英与口头输出</h3><p class="zh">先用本集词汇翻译，再用自己的话复述本集内容。</p><ol>'+''.join(items)+'</ol></div>'

def lesson_html(e):
    words=sum(c['type']=='word' for c in e['cards']); phrases=len(e['cards'])-words
    key=h(' '.join([e['id'],e['title'],e['synopsis']]).lower(),quote=True)
    return f'''<details class="episode" id="{e['id']}" data-episode="{e['id']}" data-season="{e['season']}" data-search="{key}">
<summary><span><b>{e['id']}</b>{h(e['title'])}</span><span class="episode-meta">{len(e['cards'])} 项 · {words} 词 / {phrases} 组 <span class="completed-indicator"></span></span></summary>
<div class="episode-body"><div class="episode-actions">
<a download href="downloads/episodes/{e['id']}-study-notes.txt">下载本集学习笔记 TXT</a>
<a href="{h(e['source'])}" target="_blank" rel="noopener noreferrer">字幕来源（第 {e['season']} 季合集）</a>
<label><input type="checkbox" class="done" data-id="{e['id']}"> 本集已复习</label></div>
<section class="synopsis" aria-labelledby="synopsis-{e['id']}"><div class="synopsis-heading"><h3 id="synopsis-{e['id']}">英文内容梗概 <span>含剧透</span></h3><span>{e['word_count']} / 200 words</span></div><p lang="en">{h(e['synopsis'])}</p></section>
<div class="cards">{''.join(card_html(c,i+1) for i,c in enumerate(e['cards']))}</div>{practice_html(e)}</div></details>'''

season_html=[]
for season in range(1,7):
    eps=[e for e in LESSONS if e['season']==season]
    season_html.append(f'<section class="season" id="s{season}"><div class="season-heading"><h2>第 {season} 季</h2><span>{len(eps)} 集 · {sum(len(e["cards"]) for e in eps)} 项</span></div>'+''.join(lesson_html(e) for e in eps)+'</section>')
season_options=''.join(f'<option value="{s}">第 {s} 季</option>' for s in range(1,7))
nav=''.join(f'<a href="#s{s}">第 {s} 季</a>' for s in range(1,7))
resources=['<li><a download href="downloads/study-materials.zip">全部学习资料 ZIP（含 63 集笔记与梗概）</a></li>']
for season in range(1,7):
    source=next(e['source'] for e in LESSONS if e['season']==season)
    resources.append(f'<li><a download href="downloads/season-{season:02}-study-guide.docx">第 {season} 季 Word 学习册</a> · <a href="{source}" target="_blank" rel="noopener noreferrer">双语字幕原始来源</a></li>')
refs=''.join(f'<li><a href="{r["url"]}" target="_blank" rel="noopener noreferrer">{h(r["title"])}</a></li>' for r in D['references'])
page=f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{TITLE} · 美式音标、英文释义与剧情梗概</title><meta name="description" content="风骚律师六季63集英语学习：关键词配美式IPA音标、词性与英文解释，每集不超过15项、200词以内英文梗概、{S['original_examples']}个原创双语例句，支持逐集笔记下载与复习。">
<link rel="stylesheet" href="assets/style.css?v={D['version']}"><link rel="icon" href="assets/favicon.svg" type="image/svg+xml"></head>
<body><main><header><p class="eyebrow">BETTER CALL SAUL · ENGLISH STUDY GUIDE · {VERSION}</p>
<h1>{TITLE}<small>中高级词汇 · 法律与商务语境 · 英文剧情梗概</small></h1>
<p class="intro">六季 63 集，去掉基础口语词组，聚焦关键词、专业术语与值得掌握的惯用表达。每集精选 {S['min_entries']}–{S['max_entries']} 项，单词与词组合计不超过 15 项；每项配 3 个原创双语例句。</p>
<p class="intro-detail">{S['unique_words']} 个关键词含美式 IPA 音标；全部 {S['unique_terms']} 个词条均有词性与英文释义，遮住中文后也能独立学习。</p>
<div class="stats"><div><b>63</b><span>集英文梗概 · 每集 ≤200 词</span></div><div><b>{S['entries']}</b><span>条逐集词汇记录</span></div><div><b>{S['unique_terms']}</b><span>个独立词条 · {S['unique_words']} 词 / {S['unique_phrases']} 组</span></div><div><b>{S['original_examples']}</b><span>个原创双语例句</span></div></div>
<nav>{nav}<a href="#resources">下载与使用</a><a href="{REPO}">GitHub 项目</a></nav></header>
<details class="method"><summary>学习方法与选词标准</summary><ol><li>先读本集英文梗概，再回看片段。梗概含剧透，最长 {S['max_synopsis_words']} 词。</li><li>按语境价值和难度选词，不为凑数量加入基础词组。词条显示参考字幕中的形式和回看时间，原型与实际形式可能不同。</li><li>{h(D['pronunciation_note'])}</li><li>词性提示：noun 名词，verb 动词，adjective 形容词，adverb 副词，plural noun 复数名词，phrase 词组，idiom 惯用表达，past tense / past participle 过去式 / 过去分词。</li><li>遮住中文，先读英文释义，再用 3 个原创例句练习迁移。跨集复现词条沿用例句，方便间隔复习。</li><li>完成汉译英后，用本集 3 个词条口头复述。回看时间依参考字幕版本，可能与片源有偏移。</li></ol></details>
<div class="toolbar" aria-label="筛选学习内容"><select id="season" aria-label="选择季"><option value="all">全部六季</option>{season_options}</select>
<select id="kind" aria-label="词条类型"><option value="all">单词 + 词组</option><option value="word">只看关键词</option><option value="phrase">只看词组 / 术语</option></select>
<input type="search" id="search" aria-label="搜索" placeholder="搜索词汇、英文释义、例句或剧情">
<label><input type="checkbox" id="hidezh"> 遮住中文</label><label><input type="checkbox" id="unfinished"> 仅未复习</label>
<button id="expand" type="button">展开</button><button id="collapse" type="button">收起</button><button id="print" type="button">打印</button><span id="progress"></span><span id="results" aria-live="polite"></span></div>
{''.join(season_html)}
<section class="resources" id="resources"><h2>下载与使用</h2><div class="resources-grid"><div><h3>逐集笔记与整季学习册</h3><ul>{''.join(resources)}</ul>
<p class="note">每集上方可直接下载独立 TXT 学习笔记，包含英文梗概、美式音标、词性、英文释义、用法和双语例句。公开下载不包含完整原剧台词或字幕；各季字幕来源入口指向原始合集页面。</p></div>
<div><h3>复习与导入</h3><ul><li><a download href="downloads/anki-expressions.txt">Anki 重点单词与词组（{S['unique_terms']} 条）</a></li><li><a download href="downloads/episode-expression-index.csv">逐集词汇索引 CSV</a></li><li><a download href="downloads/episode-synopses.csv">63 集英文梗概 CSV</a></li></ul>
<p>Anki 文件为 UTF-8、Tab 分隔，前两列映射正面与背面，第三列为标签，允许 HTML。同一词条只导入一次，通过逐集标签复习。</p>
<p>复习标记保存在当前浏览器中，旧版进度继续沿用。换浏览器或清除网站数据后不会保留。</p></div></div><h3>术语与使用参考</h3><ul>{refs}</ul></section>
<p class="note">美式发音核对参考 CMU Pronouncing Dictionary 与所列词典；英文释义为本项目编写。<a href="data/CMUDICT-LICENSE.txt">CMU 词典许可与致谢</a></p>
<p class="footer">更新：{UPDATED} · {VERSION}。本项目的词义说明、例句、练习与梗概为编写的学习内容；参考字幕用于选词与定位，未逐句对照音轨校订。非官方英语学习项目。</p></main><script src="assets/app.js?v={D['version']}"></script></body></html>'''
(ROOT/'index.html').write_text(page+'\n')

# Original per-episode notes; no raw screenplay or subtitle dialogue.
for e in LESSONS:
    lines=[f"{TITLE} · {e['id']} {e['title']}",f"{VERSION} · 更新 {D['updated']}",f"英文梗概（{e['word_count']} words，含剧透）",e['synopsis'],'',f"重点词汇：{len(e['cards'])} 项（单词与词组合计）",D['pronunciation_note']]
    for i,c in enumerate(e['cards'],1):
        lines.extend(['',f"{i:02}. {c['term']} · {KINDS[c['type']]} · 回看 {c['time']}"]+lexical_lines(c)+[c['meaning'],f"字幕中出现：{c['matched_form']}",c['usage']])
        if c['context']:lines.append('本集用法：'+c['context'])
        for ex in c['examples']:lines.extend([ex['en'],ex['zh']])
    lines.extend(['','本项目原创学习笔记，不含完整原剧台词。','字幕来源（整季合集）：'+e['source'],'发音核对参考：CMU Pronouncing Dictionary；https://github.com/cmusphinx/cmudict','CMU 许可：'+SITE+'data/CMUDICT-LICENSE.txt'])
    (DL/f"episodes/{e['id']}-study-notes.txt").write_text('\n'.join(lines)+'\n')

# All exports share the same data as the webpage.
with (DL/'episode-expression-index.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f,lineterminator="\n");w.writerow(['集数','标题','类型','词条','美式音标 IPA','词性','英文释义','原形','中文释义','用法','本集用法','回看时间','字幕形式'])
    for e in LESSONS:
        for c in e['cards']:w.writerow([e['id'],e['title'],KINDS[c['type']],c['term'],c['ipa'],c['part_of_speech'],c['definition_en'],c['base_form'],c['meaning'],c['usage'],c['context'],c['time'],c['matched_form']])
with (DL/'episode-synopses.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f,lineterminator="\n");w.writerow(['集数','标题','英文梗概','词数'])
    for e in LESSONS:w.writerow([e['id'],e['title'],e['synopsis'],e['word_count']])
terms={}; appearances=defaultdict(list); senses=defaultdict(lambda:defaultdict(list))
for e in LESSONS:
    for c in e['cards']:
        terms[c['term']]=c;appearances[c['term']].append(e['id'])
        senses[c['term']][(c['part_of_speech'],c['definition_en'])].append(e['id'])
anki=['#separator:Tab','#html:true','#columns:Front\tBack\tTags']
for term,c in sorted(terms.items(),key=lambda x:x[0].lower()):
    pronunciation=('<b>美式 IPA</b> '+h(c['ipa'])+'<br>') if c['ipa'] else ''
    base=('<b>原形</b> '+h(c['base_form'])+'<br>') if c['base_form'] else ''
    definitions=''.join('<p><b>'+h(pos)+'</b><br>'+h(definition)+(('<br><small>'+', '.join(ids)+'</small>') if len(senses[term])>1 else '')+'</p>' for (pos,definition),ids in senses[term].items())
    back=pronunciation+base+definitions+h(c['meaning'])+'<br>'+h(c['usage'])+'<ol>'+''.join('<li>'+h(ex['en'])+'<br>'+h(ex['zh'])+'</li>' for ex in c['examples'])+'</ol><small>复现：'+', '.join(appearances[term])+'</small>'
    tags=' '.join(['BCS',c['type']]+['BCS::'+eid for eid in appearances[term]])
    anki.append('\t'.join([h(term),back,tags]))
(DL/'anki-expressions.txt').write_text('\n'.join(anki)+'\n')

for season in range(1,7):
    doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.add_heading(f'{TITLE} · 第 {season} 季',0)
    doc.add_paragraph(VERSION+' · 每集不超过 15 项关键词与词组；关键词含美式 IPA，所有词条含词性与英文释义；每项 3 个原创双语例句。英文梗概不超过 200 词，含剧透。')
    doc.add_paragraph(D['pronunciation_note'])
    doc.add_paragraph('本文档为本项目编写的学习笔记，不含完整原剧台词或字幕。回看时间依参考字幕版本。')
    for e in [e for e in LESSONS if e['season']==season]:
        doc.add_heading(e['id']+' · '+e['title'],1)
        doc.add_heading(f"英文内容梗概（{e['word_count']} words）",2);doc.add_paragraph(e['synopsis'])
        doc.add_paragraph('双语字幕原始来源（整季合集）：'+e['source'])
        for i,c in enumerate(e['cards'],1):
            doc.add_heading(f"{i:02}. {c['term']} · {KINDS[c['type']]}",2)
            for line in lexical_lines(c):doc.add_paragraph(line)
            doc.add_paragraph(c['meaning']);doc.add_paragraph(f"回看 {c['time']} · 字幕中出现：{c['matched_form']}");doc.add_paragraph(c['usage'])
            if c['context']:doc.add_paragraph('本集用法：'+c['context'])
            for j,ex in enumerate(c['examples'],1):doc.add_paragraph(f"{j}. {ex['en']}\n{ex['zh']}")
        doc.add_heading('汉译英与复述',2)
        chosen=[e['cards'][i] for i in [0,len(e['cards'])//2,len(e['cards'])-1]]
        for i,c in enumerate(chosen,1):doc.add_paragraph(f"{i}. {c['examples'][1]['zh']}")
        doc.add_paragraph('参考答案：')
        for i,c in enumerate(chosen,1):doc.add_paragraph(f"{i}. {c['examples'][1]['en']} [{c['term']}]")
        doc.add_page_break()
    doc.add_heading('参考来源',1)
    for ref in D['references']:doc.add_paragraph(ref['title']+'：'+ref['url'])
    doc.add_paragraph('美式发音核对参考 CMU Pronouncing Dictionary；许可与致谢：'+SITE+'data/CMUDICT-LICENSE.txt')
    doc.save(DL/f'season-{season:02}-study-guide.docx')

(ROOT/'README.md').write_text(f'''# {TITLE}

网站：{SITE}

{VERSION}（{D['updated']}）覆盖六季 63 集：每集 {S['min_entries']}–15 个单词与词组，共 {S['entries']} 条逐集记录、{S['unique_terms']} 个独立词条（{S['unique_words']} 个单词、{S['unique_phrases']} 个词组或术语）、{S['original_examples']} 个原创双语例句。
全部 {S['unique_words']} 个关键词含美式 IPA；全部 {S['unique_terms']} 个词条含词性及原创英文释义。变形词补充原形，特定集数的多义词按本集用法说明。
每集有一篇英文梗概，均不超过 200 词，包含剧透。筛选侧重中高级惯用表达及法律、商务、调查语境中的关键词，难度判断不是正式 CEFR 评级。

## 功能

- 按季、关键词或词组筛选，搜索词汇、英文释义、中文、例句和剧情。
- 隐藏中文、练习汉译英、标记逐集复习进度。
- 下载 63 份逐集 TXT 学习笔记、六季 Word 学习册、Anki、词汇 CSV、梗概 CSV 及 ZIP。
- 复习进度保存在当前浏览器 localStorage，沿用旧版 `bcs-learning-v1` 键。

公开仓库和下载资料仅包含本项目编写的学习内容。字幕入口指向原始来源的整季合集，不提供完整原剧台词的公开下载。
参考字幕用于选词和定位；时间可能随片源偏移，未逐句对照音轨校订。本项目为非官方英语学习项目。

## 发音说明与致谢

{D['pronunciation_note']}
发音核对参考 [CMU Pronouncing Dictionary](https://github.com/cmusphinx/cmudict) 与页面所列 Merriam-Webster 词条。CMU 词典版权归 Carnegie Mellon University，许可见 [data/CMUDICT-LICENSE.txt](data/CMUDICT-LICENSE.txt)。本项目采用人工整理的 IPA，不逐字复制词典英文释义。
网页、逐集 TXT、Word、Anki 和 CSV 从同一份数据生成。Anki 中跨集多义词保留不同释义及对应集数。

## 预览与生成

直接打开 `index.html`，或运行 `python3 -m http.server 8766 --bind 127.0.0.1`。
网站无运行时依赖。`data/study.json` 是学习内容源；安装 `python-docx` 后，运行 `python3 tools/build.py` 可重新生成网页与下载资料。

## GitHub Pages

Settings → Pages → Deploy from a branch → `main` → `/(root)`。`.nojekyll` 使 GitHub 直接发布静态文件。
''')

bundle=DL/'study-materials.zip'
with zipfile.ZipFile(bundle,'w',zipfile.ZIP_DEFLATED) as z:
    files=[ROOT/'index.html',ROOT/'README.md',ROOT/'.nojekyll']
    for directory in ['assets','data','downloads']:
        files.extend(p for p in (ROOT/directory).rglob('*') if p.is_file() and p!=bundle and p.name!='.DS_Store')
    for p in sorted(files):z.write(p,p.relative_to(ROOT))
print(json.dumps({'episodes':len(LESSONS),'entries':S['entries'],'unique_terms':len(terms),'bundle_bytes':bundle.stat().st_size},ensure_ascii=False))
